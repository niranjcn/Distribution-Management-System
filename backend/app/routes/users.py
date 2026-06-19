import csv
import io
import logging
from fastapi import APIRouter, HTTPException, status, Depends, Query, Request, File, UploadFile, Form
from typing import Optional, List, Any, Iterable

from app.database import get_db
from app.models.user import UserCreate, UserUpdate
from app.services import user_service, reassignment_request_service, notification_service
from app.middleware.auth_middleware import get_current_user, require_admin_or_manager_or_md
from app.core.audit import audit_logger
from app.core.activity_logger import build_field_change_summary, log_business_activity
from app.utils.roles import (
    SUPER_ADMIN,
    MD_DIRECTOR,
    MANAGER,
    PDIC_STAFF,
    SUB_DISTRIBUTION_MANAGER,
    SUB_DISTRIBUTOR,
    CLUSTER,
    OPERATOR,
    normalize_role,
    can_manage_user,
    can_mutate_super_admin,
)

router = APIRouter()

logger = logging.getLogger(__name__)


def _chunks(values: List[Any], chunk_size: int) -> Iterable[List[Any]]:
    for i in range(0, len(values), chunk_size):
        yield values[i:i + chunk_size]


ALLOWED_CREATE_BY_ROLE = {
    SUPER_ADMIN: [SUPER_ADMIN, MD_DIRECTOR, MANAGER, PDIC_STAFF, SUB_DISTRIBUTION_MANAGER, SUB_DISTRIBUTOR, CLUSTER, OPERATOR],
    MANAGER: [PDIC_STAFF, SUB_DISTRIBUTION_MANAGER, SUB_DISTRIBUTOR, CLUSTER, OPERATOR],
    SUB_DISTRIBUTOR: [SUB_DISTRIBUTION_MANAGER, CLUSTER, OPERATOR],
    CLUSTER: [OPERATOR],
}


async def _branch_contains_user(root_user_id: str, target_user_id: str) -> bool:
    if str(root_user_id) == str(target_user_id):
        return True

    pending = [int(root_user_id)]
    visited = set()

    async with get_db() as db:
        while pending:
            parent_id = pending.pop()
            if parent_id in visited:
                continue
            visited.add(parent_id)

            cursor = await db.execute("SELECT id FROM users WHERE parent_id = ?", (parent_id,))
            children = await cursor.fetchall()
            for child in children:
                child_id = int(child["id"])
                if str(child_id) == str(target_user_id):
                    return True
                pending.append(child_id)

    return False


async def _can_access_user(current_user: dict, target_user: dict, *, write: bool) -> bool:
    actor_role = normalize_role(current_user.get("role"))
    target_role = normalize_role(target_user.get("role"))

    if actor_role == MD_DIRECTOR and target_role == SUPER_ADMIN:
        return False

    if str(current_user.get("id")) == str(target_user.get("id")):
        return True

    if actor_role == SUPER_ADMIN:
        if write:
            return can_mutate_super_admin(current_user.get("id"), actor_role, target_user.get("id"), target_role)
        return True

    if actor_role == MD_DIRECTOR:
        return not write

    if actor_role == PDIC_STAFF:
        return False

    if actor_role == MANAGER:
        if not write and target_role == MANAGER:
            if current_user.get("parent_id"):
                return await _branch_contains_user(current_user["parent_id"], target_user.get("id"))
            return True
        if not can_manage_user(actor_role, target_role):
            return False
        if current_user.get("parent_id"):
            return await _branch_contains_user(current_user["parent_id"], target_user.get("id"))
        return True

    if actor_role == SUB_DISTRIBUTION_MANAGER:
        if write and target_role in {SUB_DISTRIBUTOR, SUB_DISTRIBUTION_MANAGER}:
            return False
        if target_role not in {SUB_DISTRIBUTOR, SUB_DISTRIBUTION_MANAGER, CLUSTER, OPERATOR}:
            return False
        root_id = str(current_user.get("parent_id") or current_user.get("id"))
        if str(target_user.get("id")) == root_id:
            return True
        return await _branch_contains_user(root_id, target_user.get("id"))

    if actor_role == SUB_DISTRIBUTOR:
        if write:
            return False
        if target_role not in {CLUSTER, OPERATOR}:
            return False
        return await _branch_contains_user(current_user.get("id"), target_user.get("id"))

    if actor_role == CLUSTER:
        if write:
            return False
        return target_role == OPERATOR and str(target_user.get("parent_id")) == str(current_user.get("id"))

    return False


@router.get("")
async def get_users(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1),
    role: Optional[str] = None,
    status_filter: Optional[str] = Query(None, alias="status"),
    search: Optional[str] = None,
    search_by: Optional[str] = Query("all"),
    parent_id: Optional[str] = None,
    current_user: dict = Depends(get_current_user),
):
    actor_role = normalize_role(current_user.get("role"))
    normalized_role_filter = normalize_role(role) if role else None

    parent_id_filter = None
    parent_ids_in_filter = None
    roles_in_filter = None

    if actor_role in {SUPER_ADMIN, MD_DIRECTOR, MANAGER}:
        parent_id_filter = parent_id
    elif actor_role == PDIC_STAFF:
        # Staff needs role-filtered lookup for distribution recipient selectors.
        # Keep legacy self-only response for generic, non-role-filtered listing.
        if normalized_role_filter in {SUB_DISTRIBUTOR, CLUSTER, OPERATOR}:
            parent_id_filter = parent_id
        else:
            user = await user_service.get_user_by_id(str(current_user["id"]))
            return {
                "success": True,
                "message": "Users retrieved successfully",
                "data": [user] if user else [],
                "pagination": {
                    "page": page,
                    "page_size": page_size,
                    "total": 1 if user else 0,
                    "total_pages": 1,
                    "has_next": False,
                    "has_prev": False,
                },
            }
    elif actor_role == SUB_DISTRIBUTION_MANAGER:
        scope_root = str(current_user.get("parent_id") or current_user["id"])
        parent_id_filter = scope_root
        if normalized_role_filter in {CLUSTER, OPERATOR} or normalized_role_filter is None:
            manager_result = await user_service.get_users(role=SUB_DISTRIBUTION_MANAGER, parent_id=scope_root, page_size=1_000_000)
            manager_ids = [int(m["id"]) for m in manager_result["data"]]
            candidate_parent_ids = [int(scope_root)] + manager_ids if str(scope_root).isdigit() else manager_ids
            candidate_parent_ids = list(dict.fromkeys(candidate_parent_ids))

            cluster_parent_scope = candidate_parent_ids
            operator_parent_scope = candidate_parent_ids
            cluster_ids = []
            if cluster_parent_scope:
                clusters_result = await user_service.get_users(
                    role=CLUSTER,
                    parent_ids_in=cluster_parent_scope,
                    page_size=1_000_000,
                )
                cluster_ids = [int(c["id"]) for c in clusters_result["data"]]
                operator_parent_scope = list(dict.fromkeys(cluster_ids + candidate_parent_ids))

            if normalized_role_filter == CLUSTER:
                parent_ids_in_filter = candidate_parent_ids
                parent_id_filter = None
            elif normalized_role_filter == OPERATOR:
                parent_ids_in_filter = operator_parent_scope
                parent_id_filter = None
            else:
                # Read-only default: sub distribution managers see only clusters/operators in their sub-distribution branch.
                roles_in_filter = [CLUSTER, OPERATOR]
                parent_ids_in_filter = operator_parent_scope
                parent_id_filter = None
    elif actor_role == SUB_DISTRIBUTOR:
        parent_id_filter = str(current_user["id"])
        if normalized_role_filter == OPERATOR:
            sub_dist_manager_result = await user_service.get_users(role=SUB_DISTRIBUTION_MANAGER, parent_id=str(current_user["id"]), page_size=1_000_000)
            sub_dist_manager_ids = [int(m["id"]) for m in sub_dist_manager_result["data"]]
            candidate_cluster_parent_ids = [int(current_user["id"])] + sub_dist_manager_ids
            clusters_result = await user_service.get_users(role=CLUSTER, parent_ids_in=candidate_cluster_parent_ids, page_size=1_000_000)
            parent_ids_in_filter = [int(c["id"]) for c in clusters_result["data"]]
            parent_id_filter = None
    elif actor_role == CLUSTER:
        parent_id_filter = str(current_user["id"])
    elif actor_role == OPERATOR:
        if normalized_role_filter == OPERATOR:
            parent_id_filter = str(current_user.get("parent_id", ""))
        else:
            raise HTTPException(status_code=403, detail="Operators can only list operators")

    try:
        result = await user_service.get_users(
            page=page,
            page_size=page_size,
            role=normalized_role_filter,
            roles_in=roles_in_filter,
            status=status_filter,
            search=search,
            search_by=search_by,
            parent_id=parent_id_filter,
            parent_ids_in=parent_ids_in_filter,
        )

        if actor_role in {MANAGER, SUB_DISTRIBUTION_MANAGER, SUB_DISTRIBUTOR, CLUSTER}:
            filtered = []
            for row in result["data"]:
                if await _can_access_user(current_user, row, write=False):
                    filtered.append(row)
            result["data"] = filtered

        if actor_role == MD_DIRECTOR:
            result["data"] = [row for row in result["data"] if normalize_role(row.get("role")) != SUPER_ADMIN]

        return {
            "success": True,
            "message": "Users retrieved successfully",
            "data": result["data"],
            "pagination": result["pagination"],
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Unhandled route exception")
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")


@router.get("/{user_id}")
async def get_user(user_id: str, current_user: dict = Depends(get_current_user)):
    try:
        user = await user_service.get_user_by_id(user_id)
        if not user:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

        if not await _can_access_user(current_user, user, write=False):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

        return {"success": True, "message": "User retrieved successfully", "data": user}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Unhandled route exception")
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")


@router.post("", status_code=status.HTTP_201_CREATED)
async def create_user(user_data: UserCreate, current_user: dict = Depends(get_current_user)):
    actor_role = normalize_role(current_user.get("role"))
    target_role = normalize_role(user_data.role.value)

    if target_role != SUB_DISTRIBUTOR:
        user_data = user_data.model_copy(update={"digital_id": None, "broadband_id": None})

    if actor_role not in ALLOWED_CREATE_BY_ROLE:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You do not have permission to create users")

    if target_role not in ALLOWED_CREATE_BY_ROLE[actor_role]:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail=f"You cannot create role '{target_role}'")

    if target_role == SUB_DISTRIBUTION_MANAGER:
        if actor_role == SUB_DISTRIBUTOR and not user_data.parent_id:
            user_data = user_data.model_copy(update={"parent_id": str(current_user["id"])})

        if not user_data.parent_id:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Must select a sub distributor parent for sub distribution manager",
            )

        parent_user = await user_service.get_user_by_id(user_data.parent_id)
        if not parent_user or normalize_role(parent_user.get("role")) != SUB_DISTRIBUTOR:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid sub distributor selected")

        if actor_role == SUB_DISTRIBUTOR and str(user_data.parent_id) != str(current_user.get("id")):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="You can only assign sub distribution managers under your own account",
            )

    if target_role == CLUSTER:
        if actor_role == SUB_DISTRIBUTION_MANAGER and not user_data.parent_id:
            user_data = user_data.model_copy(update={"parent_id": str(current_user.get("parent_id"))})

        if not user_data.parent_id:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Must select a sub distributor parent for cluster")

        parent_user = await user_service.get_user_by_id(user_data.parent_id)
        if not parent_user or normalize_role(parent_user.get("role")) != SUB_DISTRIBUTOR:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid sub distributor selected")

        if actor_role == SUB_DISTRIBUTION_MANAGER and str(user_data.parent_id) != str(current_user.get("parent_id")):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You can only assign clusters under your own sub distribution")

        if actor_role == SUB_DISTRIBUTOR and not await _branch_contains_user(current_user.get("id"), user_data.parent_id):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Selected sub distributor is outside your branch")

    if actor_role == CLUSTER and target_role == OPERATOR and not user_data.parent_id:
        user_data = user_data.model_copy(update={"parent_id": str(current_user["id"])})

    if target_role == OPERATOR:
        if not user_data.parent_id:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Must select a cluster parent for operator")
        cluster = await user_service.get_user_by_id(user_data.parent_id)
        if not cluster or normalize_role(cluster.get("role")) != CLUSTER:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid cluster selected")

    if actor_role == SUB_DISTRIBUTOR:
        if target_role == OPERATOR:
            cluster = await user_service.get_user_by_id(user_data.parent_id)
            if not cluster or not await _branch_contains_user(current_user.get("id"), cluster.get("id")):
                raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Selected cluster is outside your branch")

    if actor_role == SUB_DISTRIBUTION_MANAGER and target_role == OPERATOR:
        cluster = await user_service.get_user_by_id(user_data.parent_id)
        if not cluster or not await _branch_contains_user(current_user.get("id"), cluster.get("id")):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Selected cluster is outside your branch")

    try:
        user = await user_service.create_user(user_data, creator_role=actor_role)
        actor_name = current_user.get("name") or current_user.get("email") or "User"
        await log_business_activity(
            user=current_user,
            path="/activity/users/create",
            description=(
                f"{actor_name} created user {user.get('name') or user.get('email') or user.get('id')} "
                f"({user.get('role') or 'unknown role'})"
            ),
        )
        return {"success": True, "message": "User created successfully", "data": user}
    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception as e:
        logger.exception("Unhandled route exception")
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")


@router.put("/{user_id}")
async def update_user(user_id: str, user_data: UserUpdate, current_user: dict = Depends(get_current_user)):
    try:
        actor_role = normalize_role(current_user.get("role"))
        if actor_role in {MD_DIRECTOR, SUB_DISTRIBUTION_MANAGER}:
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN,
                detail="This role has read-only access to users",
            )

        target_user = await user_service.get_user_by_id(user_id)
        if not target_user:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

        if not await _can_access_user(current_user, target_user, write=True):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

        if actor_role in {MD_DIRECTOR, PDIC_STAFF} and str(current_user.get("id")) != str(user_id):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

        if actor_role in {MD_DIRECTOR, PDIC_STAFF} and user_data.status is not None:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="You cannot change account status")

        user = await user_service.update_user(user_id, user_data)
        actor_name = current_user.get("name") or current_user.get("email") or "User"
        edited_fields = list(user_data.model_dump(exclude_unset=True).keys())
        change_summary = build_field_change_summary(
            before=target_user or {},
            after=user or {},
            fields=edited_fields,
            exclude_fields={"updated_at", "password_hash"},
        )
        await log_business_activity(
            user=current_user,
            path="/activity/users/update",
            description=(
                f"{actor_name} updated user "
                f"{user.get('name') or user.get('email') or user_id}; "
                f"changes: {change_summary}"
            ),
        )
        return {"success": True, "message": "User updated successfully", "data": user}
    except HTTPException:
        raise
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    except Exception as e:
        logger.exception("Unhandled route exception")
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")


async def _get_children_to_reassign(target_user: dict) -> list:
    """Get direct children that would be orphaned by deleting this user.
    Only returns immediate children — grandchildren (operators under clusters)
    stay connected to their parent automatically and don't need reassignment."""
    target_role = normalize_role(target_user.get("role"))
    children = []

    async with get_db() as db:
        if target_role == SUB_DISTRIBUTOR:
            cursor = await db.execute(
                "SELECT id, name, email, role, parent_id FROM users WHERE parent_id = ? AND role IN (?, ?, ?)",
                (int(target_user["id"]), SUB_DISTRIBUTION_MANAGER, CLUSTER, OPERATOR)
            )
            for child in await cursor.fetchall():
                children.append(dict(child))
        elif target_role == CLUSTER:
            cursor = await db.execute(
                "SELECT id, name, email, role, parent_id FROM users WHERE parent_id = ? AND role = ?",
                (int(target_user["id"]), OPERATOR)
            )
            for child in await cursor.fetchall():
                children.append(dict(child))

    return children


def _count_total_children(children: list) -> int:
    """Recursively count all items including nested ones."""
    count = 0
    for c in children:
        count += 1
        if c.get("children"):
            count += _count_total_children(c["children"])
    return count


@router.delete("/{user_id}")
async def delete_user(request: Request, user_id: str, current_user: dict = Depends(get_current_user)):
    actor_role = normalize_role(current_user.get("role"))

    if actor_role not in {SUPER_ADMIN}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

    if str(current_user.get("id")) == str(user_id):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Cannot delete your own account")

    target_user = await user_service.get_user_by_id(user_id)
    if not target_user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

    if actor_role == SUPER_ADMIN:
        if not can_mutate_super_admin(current_user.get("id"), actor_role, target_user.get("id"), target_user.get("role")):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Cannot delete another super admin")
    else:
        if not await _can_access_user(current_user, target_user, write=True):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

    target_role = normalize_role(target_user.get("role"))

    # If deleting a sub_distributor or cluster, create reassignment request instead
    if target_role in {SUB_DISTRIBUTOR, CLUSTER}:
        children = await _get_children_to_reassign(target_user)

        if not children:
            # No children orphaned, proceed with normal deletion
            try:
                success = await user_service.delete_user(user_id)
                if not success:
                    raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

                actor_name = current_user.get("name") or current_user.get("email") or "User"
                target_name = target_user.get("name") or target_user.get("email") or user_id
                await log_business_activity(
                    user=current_user,
                    path="/activity/users/delete",
                    description=f"{actor_name} deleted user {target_name}",
                )
                audit_logger.warning(
                    "USER_DELETE | actor_id=%s | actor_email=%s | target_user_id=%s | ip=%s",
                    current_user.get("id"),
                    current_user.get("email"),
                    user_id,
                    request.client.host if request.client else "unknown",
                )
                return {"success": True, "message": "User deleted successfully"}
            except HTTPException:
                raise
            except Exception as e:
                logger.exception("Unhandled route exception")
                raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")

        # Create reassignment request
        actor_name = current_user.get("name") or current_user.get("email") or "User"
        req = await reassignment_request_service.create_reassignment_request(
            deleted_user=target_user,
            children=children,
            actor_name=actor_name,
        )

        # Notify all super admins
        super_admins = await user_service.get_users(role=SUPER_ADMIN, page_size=10000)
        for sa in super_admins["data"]:
            await notification_service.create_notification(
                user_id=sa["id"],
                title="Reassignment Request",
                message=f"A {target_role} ({target_user.get('name')}) was scheduled for deletion. {len(children)} user(s) need reassignment.",
                notification_type="warning",
                category="user",
                link="/reassignment-requests",
                metadata={"request_id": req.get("request_id"), "request_db_id": req.get("id")},
            )

        return {
            "success": True,
            "message": f"Reassignment request created. {len(children)} user(s) need to be reassigned before deletion.",
            "data": {"request": req, "children_count": len(children)},
        }

    # For non-sub_distributor/cluster, proceed with normal deletion
    try:
        success = await user_service.delete_user(user_id)
        if not success:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

        actor_name = current_user.get("name") or current_user.get("email") or "User"
        target_name = target_user.get("name") or target_user.get("email") or user_id
        await log_business_activity(
            user=current_user,
            path="/activity/users/delete",
            description=f"{actor_name} deleted user {target_name}",
        )

        audit_logger.warning(
            "USER_DELETE | actor_id=%s | actor_email=%s | target_user_id=%s | ip=%s",
            current_user.get("id"),
            current_user.get("email"),
            user_id,
            request.client.host if request.client else "unknown",
        )

        return {"success": True, "message": "User deleted successfully"}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Unhandled route exception")
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")


@router.patch("/{user_id}/status")
async def update_user_status(
    request: Request,
    user_id: str,
    status_update: dict,
    current_user: dict = Depends(get_current_user),
):
    actor_role = normalize_role(current_user.get("role"))
    if actor_role not in {SUPER_ADMIN, MANAGER}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

    status_value = status_update.get("status")
    if status_value not in ["active", "inactive", "suspended"]:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid status value")

    target_user = await user_service.get_user_by_id(user_id)
    if not target_user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

    if not await _can_access_user(current_user, target_user, write=True):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

    if actor_role in {MANAGER, SUB_DISTRIBUTION_MANAGER} and normalize_role(target_user.get("role")) == SUPER_ADMIN:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Cannot update super admin status")

    try:
        user = await user_service.update_user_status(user_id, status_value)
        actor_name = current_user.get("name") or current_user.get("email") or "User"
        change_summary = build_field_change_summary(
            before=target_user or {},
            after=user or {},
            fields=["status"],
            exclude_fields={"updated_at"},
        )
        await log_business_activity(
            user=current_user,
            path="/activity/users/status-update",
            description=(
                f"{actor_name} updated user status for "
                f"{user.get('name') or user.get('email') or user_id}; "
                f"changes: {change_summary}"
            ),
        )
        audit_logger.info(
            "USER_STATUS_UPDATE | actor_id=%s | actor_email=%s | target_user_id=%s | status=%s | ip=%s",
            current_user.get("id"),
            current_user.get("email"),
            user_id,
            status_value,
            request.client.host if request.client else "unknown",
        )
        return {"success": True, "message": "User status updated successfully", "data": user}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Unhandled route exception")
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")


@router.patch("/{user_id}/credentials")
async def admin_update_credentials(
    request: Request,
    user_id: str,
    data: dict,
    current_user: dict = Depends(get_current_user),
):
    from app.utils.security import get_password_hash as _hash
    from datetime import datetime as _dt, timezone as _timezone

    actor_role = normalize_role(current_user.get("role"))
    actor_id = int(current_user.get("id"))
    target_id = int(user_id)
    is_self = actor_id == target_id

    if actor_role != SUPER_ADMIN:
        # Non-super-admins can only update their own email, not password
        if not is_self:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")
        if "password" in data and data["password"]:
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Only super admin can change password")
        if "email" not in data or not data["email"]:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Email is required")

    target_user = await user_service.get_user_by_id(user_id)
    if not target_user:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="User not found")

    if not can_mutate_super_admin(actor_id, actor_role, target_user.get("id"), target_user.get("role")):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Cannot update another super admin credentials")

    try:
        async with get_db() as db:
            update_fields = []
            params = []

            if "email" in data and data["email"]:
                normalized_email = str(data["email"]).lower().strip()
                cursor = await db.execute("SELECT id FROM users WHERE email = ? AND id != ?", (normalized_email, target_id))
                if await cursor.fetchone():
                    raise HTTPException(status_code=400, detail="Email already in use")
                update_fields.append("email = ?")
                params.append(normalized_email)

            if "password" in data and data["password"]:
                if len(data["password"]) < 8:
                    raise HTTPException(status_code=400, detail="Password must be at least 8 characters")
                update_fields.append("password_hash = ?")
                params.append(_hash(data["password"]))

            if not update_fields:
                raise HTTPException(status_code=400, detail="No data to update")

            update_fields.append("updated_at = ?")
            params.append(_dt.now().replace(tzinfo=None).isoformat())
            params.append(int(user_id))

            cursor = await db.execute(f"UPDATE users SET {', '.join(update_fields)} WHERE id = ?", params)
            if cursor.rowcount == 0:
                raise HTTPException(status_code=404, detail="User not found")
            await db.commit()

        updated = await user_service.get_user_by_id(user_id)

        audit_logger.info(
            "USER_CREDENTIALS_UPDATE | actor_id=%s | actor_email=%s | target_user_id=%s | ip=%s",
            current_user.get("id"),
            current_user.get("email"),
            user_id,
            request.client.host if request.client else "unknown",
        )

        return {"success": True, "message": "Credentials updated", "data": updated}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Unhandled route exception")
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")


@router.get("/role/{role}")
async def get_users_by_role(role: str, current_user: dict = Depends(get_current_user)):
    actor_role = normalize_role(current_user.get("role"))
    normalized = normalize_role(role)

    if actor_role not in {SUPER_ADMIN, MD_DIRECTOR, MANAGER, SUB_DISTRIBUTION_MANAGER, SUB_DISTRIBUTOR, CLUSTER}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

    try:
        users = await user_service.get_users_by_role(normalized)
        filtered = []
        for row in users:
            if await _can_access_user(current_user, row, write=False):
                filtered.append(row)
        return {"success": True, "message": "Users retrieved successfully", "data": filtered}
    except HTTPException:
        raise
    except Exception as e:
        logger.exception("Unhandled route exception")
        raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="An internal error occurred. Please try again later.")


@router.post("/bulk-upload", status_code=status.HTTP_201_CREATED)
async def bulk_upload_users(
    request: Request,
    file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user),
):
    import asyncio

    actor_role = normalize_role(current_user.get("role"))
    if actor_role not in {SUPER_ADMIN, MANAGER}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Insufficient permissions")

    if not file.filename:
        raise HTTPException(status_code=400, detail="No file provided")

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in {"csv", "xlsx", "xls"}:
        raise HTTPException(status_code=400, detail="Unsupported file format. Use .csv, .xlsx, or .xls")

    try:
        contents = await file.read()
        rows = _parse_file(contents, ext)
    except Exception as e:
        logger.exception("File parse error")
        raise HTTPException(status_code=400, detail=f"Failed to parse file: {str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="File is empty or has no data rows")

    from app.utils.security import get_password_hash as _hash
    from datetime import datetime as _dt, timezone as _tz

    actor_name = current_user.get("name") or current_user.get("email") or "User"

    async def _log_bulk_upload_summary(created_count: int, skipped_count: int, error_count: int) -> None:
        await log_business_activity(
            user=current_user,
            path="/activity/users/bulk-upload",
            description=(
                f"{actor_name} used bulk upload for users: "
                f"{created_count} created, {skipped_count} skipped, {error_count} errors"
            ),
        )

    created = []
    errors = []
    skipped = []
    prepared_rows = []
    seen_emails = set()

    for idx, row in enumerate(rows):
        row_num = idx + 2
        role_val = str(row.get("role") or "").strip().lower()
        email = str(row.get("email") or "").strip().lower()
        password = str(row.get("password") or "")
        name = str(row.get("name") or "").strip()
        digital_id = str(row.get("digital_id") or "").strip() or None
        broadband_id = str(row.get("broadband_id") or "").strip() or None
        phone = str(row.get("phone") or "").strip() or None
        department = str(row.get("department") or "").strip() or None
        location = str(row.get("location") or "").strip() or None

        if not email or not password or not name:
            errors.append({"row": row_num, "email": email, "error": "Missing required fields (email, password, name)"})
            continue
        if not role_val:
            errors.append({"row": row_num, "email": email, "error": "Missing role"})
            continue

        normalized_role = normalize_role(role_val)
        if not normalized_role or normalized_role not in {"sub_distributor", "sub_distribution_manager", "cluster", "operator"}:
            errors.append({"row": row_num, "email": email, "error": f"Invalid role '{role_val}'. Allowed: sub_distributor, sub_distribution_manager, cluster, operator"})
            continue

        if email in seen_emails:
            skipped.append({"row": row_num, "email": email, "reason": "Duplicate email in file"})
            continue
        seen_emails.add(email)

        sd_email = str(row.get("sub_distributor_email") or "").strip().lower() or None
        cluster_email = str(row.get("cluster_email") or "").strip().lower() or None

        prepared_rows.append({
            "row": row_num,
            "email": email,
            "password": password,
            "name": name,
            "normalized_role": normalized_role,
            "digital_id": digital_id,
            "broadband_id": broadband_id,
            "phone": phone,
            "department": department,
            "location": location,
            "sd_email": sd_email,
            "cluster_email": cluster_email,
        })

    # Pre-hash all passwords before DB connection (bcrypt is CPU-intensive, ~50ms each).
    # Run in parallel via gather to avoid sequential ~50ms × N delay.
    loop = asyncio.get_running_loop()
    hashed = await asyncio.gather(
        *(loop.run_in_executor(None, _hash, item["password"]) for item in prepared_rows)
    )
    for item, pw_hash in zip(prepared_rows, hashed):
        item["password_hash"] = pw_hash

    if not prepared_rows:
        await _log_bulk_upload_summary(0, len(skipped), len(errors))
        return {
            "success": True,
            "message": f"Bulk upload complete: 0 created, {len(skipped)} skipped, {len(errors)} errors",
            "data": {
                "created_count": 0,
                "skipped_count": len(skipped),
                "error_count": len(errors),
                "created": created,
                "skipped": skipped,
                "errors": errors,
            }
        }

    async with get_db() as db:
        all_emails = [item["email"] for item in prepared_rows]
        existing_emails = set()
        if all_emails:
            for batch in _chunks(all_emails, 500):
                placeholders = ",".join(["?"] * len(batch))
                cursor = await db.execute(
                    f"SELECT LOWER(email) as email FROM users WHERE LOWER(email) IN ({placeholders})",
                    batch,
                )
                rows_found = await cursor.fetchall()
                for r in rows_found:
                    val = r.get("email") or r.get(0)
                    if val:
                        existing_emails.add(str(val).strip().lower())

        # Batch-resolve parent references (like device bulk upload uses _fetch_existing_values)
        sd_emails = set()
        cluster_emails = set()
        for item in prepared_rows:
            if item["normalized_role"] in ("cluster", "sub_distribution_manager") and item["sd_email"]:
                sd_emails.add(item["sd_email"])
            if item["normalized_role"] == "operator" and item["cluster_email"]:
                cluster_emails.add(item["cluster_email"])

        sd_parent_map = {}
        if sd_emails:
            for batch in _chunks(list(sd_emails), 500):
                placeholders = ",".join(["?"] * len(batch))
                cursor = await db.execute(
                    f"SELECT LOWER(email) as email, id FROM users WHERE LOWER(email) IN ({placeholders}) AND role = ?",
                    [*batch, "sub_distributor"],
                )
                for row in await cursor.fetchall():
                    sd_parent_map[row["email"]] = int(row["id"])

        cluster_parent_map = {}
        if cluster_emails:
            for batch in _chunks(list(cluster_emails), 500):
                placeholders = ",".join(["?"] * len(batch))
                cursor = await db.execute(
                    f"SELECT LOWER(email) as email, id FROM users WHERE LOWER(email) IN ({placeholders}) AND role = ?",
                    [*batch, "cluster"],
                )
                for row in await cursor.fetchall():
                    cluster_parent_map[row["email"]] = int(row["id"])

        insertable_rows = []
        for item in prepared_rows:
            if item["email"] in existing_emails:
                skipped.append({"row": item["row"], "email": item["email"], "reason": "Email already exists"})
                continue

            parent_id = None
            if item["normalized_role"] in ("cluster", "sub_distribution_manager"):
                if not item["sd_email"]:
                    errors.append({"row": item["row"], "email": item["email"], "error": f"sub_distributor_email is required for role '{item['normalized_role']}'"})
                    continue
                parent_id = sd_parent_map.get(item["sd_email"])
                if parent_id is None:
                    errors.append({"row": item["row"], "email": item["email"], "error": f"Sub-distributor with email '{item['sd_email']}' not found"})
                    continue

            if item["normalized_role"] == "operator":
                if not item["cluster_email"]:
                    errors.append({"row": item["row"], "email": item["email"], "error": "cluster_email is required for role 'operator'"})
                    continue
                parent_id = cluster_parent_map.get(item["cluster_email"])
                if parent_id is None:
                    errors.append({"row": item["row"], "email": item["email"], "error": f"Cluster with email '{item['cluster_email']}' not found"})
                    continue

            item["parent_id"] = parent_id
            insertable_rows.append(item)

        if not insertable_rows:
            await _log_bulk_upload_summary(0, len(skipped), len(errors))
            return {
                "success": True,
                "message": f"Bulk upload complete: 0 created, {len(skipped)} skipped, {len(errors)} errors",
                "data": {
                    "created_count": 0,
                    "skipped_count": len(skipped),
                    "error_count": len(errors),
                    "created": created,
                    "skipped": skipped,
                    "errors": errors,
                }
            }

        insert_sql = """INSERT INTO users (email, password_hash, name, role, digital_id, broadband_id,
            phone, department, location, status, parent_id, permissions,
            theme, compact_mode, email_notifications, push_notifications,
            is_verified, created_at, updated_at, last_login)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)"""

        should_commit = True
        for batch in _chunks(insertable_rows, 500):
            batch_payload = []
            for item in batch:

                now = _dt.now(_tz.utc).replace(tzinfo=None).isoformat()
                batch_payload.append((
                    item["email"],
                    item["password_hash"],
                    item["name"],
                    item["normalized_role"],
                    item["digital_id"],
                    item["broadband_id"],
                    item["phone"],
                    item["department"],
                    item["location"],
                    "active",
                    item.get("parent_id"),
                    "{}",
                    "light",
                    0,
                    1,
                    1,
                    0,
                    now,
                    now,
                    None,
                ))

            try:
                await db.executemany(insert_sql, batch_payload)
                for item in batch:
                    created.append({"row": item["row"], "email": item["email"], "role": item["normalized_role"], "name": item["name"]})
            except Exception as batch_error:
                for item in batch:
                    row_idx = item["row"]
                    email = item["email"]
                    try:
                        now = _dt.now(_tz.utc).replace(tzinfo=None).isoformat()

                        await db.execute(
                            insert_sql,
                            (
                                email,
                                item["password_hash"],
                                item["name"],
                                item["normalized_role"],
                                item["digital_id"],
                                item["broadband_id"],
                                item["phone"],
                                item["department"],
                                item["location"],
                                "active",
                                item.get("parent_id"),
                                "{}",
                                "light",
                                0,
                                1,
                                1,
                                0,
                                now,
                                now,
                                None,
                            ),
                        )
                        created.append({"row": row_idx, "email": email, "role": item["normalized_role"], "name": item["name"]})
                    except Exception as single_error:
                        lowered = str(single_error).lower()
                        if "duplicate" in lowered or "unique" in lowered:
                            skipped.append({"row": row_idx, "email": email, "reason": "Email already exists"})
                        else:
                            errors.append({"row": row_idx, "email": email, "error": str(single_error)[:200]})
                            should_commit = False
                            break

                if not should_commit:
                    break

                logger.warning("Batch insert fallback triggered for users due to: %s", str(batch_error))

            await asyncio.sleep(0)

        if should_commit and insertable_rows:
            await db.commit()
        elif not insertable_rows:
            pass
        else:
            await db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail="Bulk upload was rolled back due to an unexpected insert error. Please retry."
            )

    await _log_bulk_upload_summary(len(created), len(skipped), len(errors))

    audit_logger.info(
        "USER_BULK_UPLOAD | actor=%s | total=%d | created=%d | skipped=%d | errors=%d",
        current_user.get("email"), len(prepared_rows), len(created), len(skipped), len(errors),
    )

    return {
        "success": True,
        "message": f"Bulk upload complete: {len(created)} created, {len(skipped)} skipped, {len(errors)} errors",
        "data": {
            "total": len(prepared_rows),
            "created_count": len(created),
            "skipped_count": len(skipped),
            "error_count": len(errors),
            "created": created,
            "skipped": skipped,
            "errors": errors,
        },
    }


def _parse_file(contents: bytes, ext: str) -> list:
    rows = []
    if ext == "csv":
        decoded = contents.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(decoded))
        for row in reader:
            rows.append({k.strip().lower(): v.strip() if v else "" for k, v in row.items()})
    elif ext in ("xlsx", "xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
        ws = wb.active
        header_row = None
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            values = [str(v).strip() if v is not None else "" for v in row]
            if r_idx == 0:
                header_row = [str(h).strip().lower() for h in values]
                continue
            if header_row:
                row_dict = {}
                for c_idx, val in enumerate(values):
                    if c_idx < len(header_row):
                        row_dict[header_row[c_idx]] = val
                if any(row_dict.values()):
                    rows.append(row_dict)
        wb.close()
    return rows



